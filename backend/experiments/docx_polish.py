from __future__ import annotations

import copy
import re
from typing import List

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

URL_RX = re.compile(r"(https?://[^\s<>\"]+?)(?=[\s\)\]]|$|[.,;]\s|[.,;]$)")
CIT_RX = re.compile(r"\[(\d{1,3}(?:\s*[,–-]\s*\d{1,3})*)\]")
NUM_RX = re.compile(r"\d{1,3}")


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def _run(text: str, rpr_template=None, bold=None, color=None, underline=None, size=None):
    r = OxmlElement("w:r")
    rpr = copy.deepcopy(rpr_template) if rpr_template is not None else OxmlElement("w:rPr")
    if bold:
        b = OxmlElement("w:b"); rpr.append(b)
    if color:
        c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    if underline:
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    return r


def _hyperlink_external(part, url: str, text: str, rpr_template=None):
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    h.append(_run(text, rpr_template))   # black, no underline: same look as body text
    return h


def _hyperlink_anchor(anchor: str, text: str, rpr_template=None):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor)
    h.append(_run(text, rpr_template))
    return h


def _first_rpr(p: Paragraph):
    r = p._p.find(".//" + qn("w:rPr"))
    if r is not None:
        r = copy.deepcopy(r)
        for tag in ("w:b", "w:color", "w:u"):
            for el in r.findall(qn(tag)):
                r.remove(el)
    return r


def _clear_content(p: Paragraph):
    for child in list(p._p):
        if not child.tag.endswith("}pPr"):
            p._p.remove(child)


def _bookmark(p: Paragraph, name: str, bid: int):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(s)
    else:
        p._p.insert(0, s)
    p._p.append(e)


def _all_paragraphs(doc) -> List[Paragraph]:
    out = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
    return out


# --------------------------------------------------------------------------- #
# references
# --------------------------------------------------------------------------- #
def make_references_clickable(doc) -> dict:
    part = doc.part
    refs = [p for p in doc.paragraphs if re.match(r"^\[\d+\]\t", p.text)]
    ref_nums = set()
    bid = 9000
    for p in refs:
        n = int(re.match(r"^\[(\d+)\]", p.text).group(1))
        ref_nums.add(n)
        text = p.text
        rpr = _first_rpr(p)
        _clear_content(p)
        pos = 0
        for m in URL_RX.finditer(text):
            if m.start() > pos:
                p._p.append(_run(text[pos:m.start()], rpr))
            p._p.append(_hyperlink_external(part, m.group(1), m.group(1), rpr))
            pos = m.end()
        if pos < len(text):
            p._p.append(_run(text[pos:], rpr))
        _bookmark(p, f"ref{n}", bid); bid += 1

    n_links = 0
    ref_elems = {id(p._p) for p in refs}      # Paragraph objects are recreated on access; compare elements
    for p in _all_paragraphs(doc):
        if id(p._p) in ref_elems or not CIT_RX.search(p.text):
            continue
        if p.style is not None and p.style.name.startswith("Heading"):
            continue
        text = p.text
        rpr = _first_rpr(p)
        _clear_content(p)
        pos = 0
        for m in CIT_RX.finditer(text):
            if m.start() > pos:
                p._p.append(_run(text[pos:m.start()], rpr))
            inner = m.group(1)
            p._p.append(_run("[", rpr))
            ipos = 0
            for nm in NUM_RX.finditer(inner):
                if nm.start() > ipos:
                    p._p.append(_run(inner[ipos:nm.start()], rpr))
                n = int(nm.group(0))
                if n in ref_nums:
                    p._p.append(_hyperlink_anchor(f"ref{n}", nm.group(0), rpr)); n_links += 1
                else:
                    p._p.append(_run(nm.group(0), rpr))
                ipos = nm.end()
            if ipos < len(inner):
                p._p.append(_run(inner[ipos:], rpr))
            p._p.append(_run("]", rpr))
            pos = m.end()
        if pos < len(text):
            p._p.append(_run(text[pos:], rpr))
    return {"references": len(refs), "citation_links": n_links}


# --------------------------------------------------------------------------- #
# captions
# --------------------------------------------------------------------------- #
def style_captions(doc) -> int:
    n = 0
    for p in doc.paragraphs:
        m = re.match(r"^((?:Fig\.|Table) \d+)\s{1,2}(.*)$", p.text, re.S)
        if not m:
            continue
        rpr = _first_rpr(p)
        _clear_content(p)
        p._p.append(_run(m.group(1), rpr, bold=True))
        p._p.append(_run("  " + m.group(2), rpr))
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        if m.group(1).startswith("Table"):
            p.paragraph_format.keep_with_next = True
        else:
            prev = p._p.getprevious()
            if prev is not None and prev.xpath(".//w:drawing"):
                Paragraph(prev, p._parent).paragraph_format.keep_with_next = True
        n += 1
    return n


# --------------------------------------------------------------------------- #
# tables
# --------------------------------------------------------------------------- #
def _set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct"); tblPr.append(w)
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "autofit"); tblPr.append(lay)
    # cell margins
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("left", "60"), ("right", "60"), ("top", "20"), ("bottom", "20")):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), val); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)


def style_tables(doc, font_pt: float = 8.0) -> int:
    for tbl in doc.tables:
        try:
            tbl.style = "Table Grid"
        except Exception:
            pass
        _set_borders(tbl)
        # remove explicit cell widths so autofit can distribute columns
        for tc in tbl._tbl.iter(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is not None:
                for w in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(w)
        for ri, row in enumerate(tbl.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(font_pt)
                        if ri == 0:
                            r.font.bold = True
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for old in tcPr.findall(qn("w:shd")):
                        tcPr.remove(old)
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E7E6E6")
                    tcPr.append(shd)
        # repeat header row on page breaks
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        if trPr.find(qn("w:tblHeader")) is None:
            h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    return len(doc.tables)


def polish(doc) -> dict:
    out = {"tables": style_tables(doc), "captions": style_captions(doc)}
    out.update(make_references_clickable(doc))
    return out
